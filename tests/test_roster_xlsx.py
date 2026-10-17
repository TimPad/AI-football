"""Synthetic workbooks: no participant data or real spreadsheets."""
from io import BytesIO
from zipfile import ZipFile
import pytest
from openpyxl import Workbook, load_workbook


def book(rows, title='Состав', second=None):
    workbook = Workbook()
    workbook.active.title = title
    for row in rows:
        workbook.active.append(row)
    if second:
        sheet = workbook.create_sheet('Вторая команда')
        for row in second:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse(value):
    from aifootball.service.catalog.roster_imports.xlsx import read_workbook
    return read_workbook(value, 'состав.xlsx')


def test_read_workbook_keeps_source_rows_types_and_sheets():
    from datetime import date
    source = parse(book([['ФИО', 'Номер', 'Дата рождения'], ['Первый Игрок', '07', date(2000, 2, 3)],
                         [None], ['Второй Игрок', 9, None]], second=[['ФИО'], ['Третий Игрок']]))
    assert len(source['sheets']) == 2
    row = source['sheets'][0]['rows'][1]
    assert row['source_row'] == 2
    assert row['cells'] == ['Первый Игрок', '07', '2000-02-03']
    assert source['sheets'][0]['rows'][2]['source_row'] == 4


def test_formula_is_preserved_as_an_issue_not_evaluated():
    source = parse(book([['ФИО', 'Номер'], ['Первый Игрок', '=2+5']]))
    assert source['sheets'][0]['rows'][1]['formulas'] == [1]


def test_exactly_1000_people_are_allowed_across_sheets():
    rows = [['ФИО']] + [[f'Игрок {i}'] for i in range(1000)]
    assert len(parse(book(rows))['sheets'][0]['rows']) == 1001


def test_rejects_unsupported_corrupt_and_excess_rows():
    from aifootball.service.catalog.roster_imports.xlsx import read_workbook
    from aifootball.service.errors import DomainError
    with pytest.raises(DomainError, match='xlsx'):
        read_workbook(b'old', 'состав.xls')
    with pytest.raises(DomainError):
        parse(b'broken')
    with pytest.raises(DomainError, match='1 000'):
        parse(book([['ФИО']] + [[f'Игрок {i}'] for i in range(1001)]))


def test_template_and_report_cells_never_execute_formulas():
    from aifootball.service.catalog.roster_imports.xlsx import template, report_workbook
    source = load_workbook(BytesIO(template('tournament')))
    assert [c.value for c in source.active[1]][:3] == ['Команда', 'ФИО', 'Номер']
    report = load_workbook(BytesIO(report_workbook([{'source_row': 2, 'name': '=1+1',
        'team_name': '+Опасный ввод', 'issues': ['Проверьте строку']}])) )
    assert all(cell.data_type != 'f' for row in report.active for cell in row)
