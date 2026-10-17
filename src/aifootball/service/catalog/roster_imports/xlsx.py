"""Bounded XLSX values for review; formulas are never evaluated."""
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from aifootball.service.errors import DomainError

FILE_LIMIT = 5 * 1024 * 1024
UNPACKED_LIMIT = 25 * 1024 * 1024
ROW_LIMIT = 1000
COLUMN_LIMIT = 64
PHYSICAL_ROW_LIMIT = 10000
HEADERS = ['Команда', 'ФИО', 'Номер', 'Амплуа', 'Дата рождения', 'Рост, см', 'Вес, кг']


def read_workbook(content: bytes, filename: str) -> dict:
    if not filename.lower().endswith('.xlsx'):
        raise DomainError('xlsx_format', 'Сохраните таблицу в формате .xlsx (Excel)')
    if len(content) > FILE_LIMIT:
        raise DomainError('xlsx_size', 'Файл Excel больше 5 МиБ')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 512 or sum(item.file_size for item in entries) > UNPACKED_LIMIT:
                raise DomainError('xlsx_size', 'Распакованная книга больше 25 МиБ')
            if any('vbaProject' in item.filename for item in entries):
                raise DomainError('xlsx_format', 'Сохраните книгу без макросов в .xlsx')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            result = []
            count = 0
            for sheet_index, sheet in enumerate(workbook.worksheets):
                if sheet.max_row and sheet.max_row > PHYSICAL_ROW_LIMIT:
                    raise DomainError('xlsx_rows', 'В книге слишком большие пропуски строк (не более 10 000 строк листа)')
                if sheet.max_column and sheet.max_column > COLUMN_LIMIT:
                    raise DomainError('xlsx_columns', 'Оставьте не более 64 столбцов в листе состава')
                sheet.reset_dimensions()
                rows = []
                for source_row, cells in enumerate(sheet.iter_rows(max_row=PHYSICAL_ROW_LIMIT, max_col=COLUMN_LIMIT), 1):
                    values, formulas = [], []
                    for column, cell in enumerate(cells):
                        value = cell.value
                        if cell.data_type == 'f':
                            formulas.append(column)
                        if isinstance(value, datetime):
                            value = value.date().isoformat()
                        elif isinstance(value, date):
                            value = value.isoformat()
                        elif value is not None and not isinstance(value, (str, int, float, bool)):
                            value = str(value)
                        values.append(value)
                    while values and values[-1] is None:
                        values.pop()
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    # The first populated row of each sheet is its header.
                    if rows:
                        count += 1
                    if count > ROW_LIMIT:
                        raise DomainError('xlsx_rows', 'В книге разрешено не более 1 000 непустых строк состава')
                    rows.append({'id': f'{sheet_index}:{source_row}', 'source_row': source_row,
                                 'cells': values, 'formulas': formulas})
                if rows:
                    result.append({'name': sheet.title, 'rows': rows})
            if not result:
                raise DomainError('xlsx_empty', 'В книге нет заполненных строк')
            return {'sheets': result}
        finally:
            workbook.close()
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError('xlsx_invalid', 'Не удалось прочитать Excel. Проверьте, что книга не повреждена и не защищена паролем.') from exc


def _export(headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Состав'
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='19221F')
        cell.fill = PatternFill('solid', fgColor='EEF8DC')
    sheet.freeze_panes = 'A2'
    for values in rows:
        sheet.append(values)
        # Every imported string is a literal, even if it begins with '='.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 25
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def template(scope: str) -> bytes:
    return _export(HEADERS if scope == 'tournament' else HEADERS[1:], [])


def report_workbook(rows: list[dict]) -> bytes:
    return _export(['Строка Excel', 'Команда', 'ФИО', 'Действие', 'Проверка'], [
        [row.get('source_row'), row.get('team_name', ''), row.get('name', ''),
         'Пропущена' if row.get('excluded') else row.get('action', ''), '; '.join(row.get('issues', []))]
        for row in rows])
